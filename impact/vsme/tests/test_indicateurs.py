from datetime import date
from io import BytesIO

import pytest
from django.conf import settings
from django.urls import reverse
from openpyxl import load_workbook
from pptx import Presentation
from pytest_django.asserts import assertTemplateUsed

from utils.pptx import PPTX_CONTENT_TYPE
from utils.xlsx import XLSX_CONTENT_TYPE
from vsme.models import EXIGENCES_DE_PUBLICATION
from vsme.models import RapportVSME

INDICATEURS_VSME_BASE_URL = "/indicateurs/vsme/"
INDICATEURS_VSME_URL = INDICATEURS_VSME_BASE_URL + "{id_public}/{annee}/"
CATEGORIE_VSME_URL = INDICATEURS_VSME_BASE_URL + "{vsme_id}/categorie/{categorie_id}/"
EXIGENCE_DE_PUBLICATION_URL = (
    INDICATEURS_VSME_BASE_URL
    + "{vsme_id}/exigence-de-publication/{exigence_de_publication_code}/"
)


def test_categories_vsme_est_prive(client, entreprise_factory, alice):
    entreprise = entreprise_factory()

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.get(url)

    assert response.status_code == 302
    connexion_url = reverse("users:login")
    assert response.url == f"{connexion_url}?next={url}"

    client.force_login(alice)
    response = client.get(url)

    assert response.status_code == 403


def test_categories_vsme_avec_utilisateur_authentifie(
    client, entreprise_factory, alice
):
    entreprise = entreprise_factory(utilisateur=alice)
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.get(url)

    assert response.status_code == 200
    assertTemplateUsed(response, "base.html")
    assertTemplateUsed(response, "snippets/tableau_de_bord_menu.html")
    assertTemplateUsed(response, "vsme/categories.html")


def test_categories_vsme_sans_siren_et_sans_annee(client, entreprise_factory, alice):
    entreprise = entreprise_factory(utilisateur=alice)
    client.force_login(alice)

    url = INDICATEURS_VSME_BASE_URL
    response = client.get(url, follow=True)

    assert response.status_code == 200
    assert response.redirect_chain == [
        (f"/indicateurs/vsme/{entreprise.id_public}/", 302)
    ]


@pytest.mark.parametrize(
    "categorie_id", ["informations-generales", "environnement", "social", "gouvernance"]
)
def test_categorie_vsme_avec_utilisateur_authentifie(
    categorie_id, client, alice, rapport_vsme
):
    client.force_login(alice)

    url = CATEGORIE_VSME_URL.format(vsme_id=rapport_vsme.id, categorie_id=categorie_id)
    response = client.get(url)

    assert response.status_code == 200
    # vérifie que tous les templates sont bien chargés dans la page résultante
    assertTemplateUsed(response, "base.html")
    assertTemplateUsed(response, "snippets/tableau_de_bord_menu.html")
    assertTemplateUsed(response, "vsme/categorie.html")
    assert response.context["categorie"].value["id"] == categorie_id


def test_categorie_vsme_inexistante_retourne_une_404(client, alice, rapport_vsme):
    client.force_login(alice)

    url = CATEGORIE_VSME_URL.format(
        vsme_id=rapport_vsme.id, categorie_id="n-existe-pas"
    )
    response = client.get(url)

    assert response.status_code == 404


def test_categorie_vsme_d_un_rapport_inexistant_retourne_une_404(client, alice):
    client.force_login(alice)

    url = CATEGORIE_VSME_URL.format(vsme_id="yolo", categorie_id="environnement")

    response = client.get(url)

    assert response.status_code == 404


def test_categorie_vsme_est_privee(client, bob, rapport_vsme):
    url = CATEGORIE_VSME_URL.format(
        vsme_id=rapport_vsme.id, categorie_id="environnement"
    )
    response = client.get(url)

    assert response.status_code == 302
    connexion_url = reverse("users:login")
    assert response.url == f"{connexion_url}?next={url}"

    # Bob n'est pas rattaché à l'entreprise du rapport VSME
    client.force_login(bob)
    response = client.get(url)

    assert response.status_code == 403


@pytest.mark.parametrize("exigence_de_publication", EXIGENCES_DE_PUBLICATION.values())
def test_exigence_de_publication_vsme_avec_utilisateur_authentifie(
    exigence_de_publication, client, alice, rapport_vsme
):
    client.force_login(alice)

    url = EXIGENCE_DE_PUBLICATION_URL.format(
        vsme_id=rapport_vsme.id,
        exigence_de_publication_code=exigence_de_publication.code,
    )
    response = client.get(url)

    assert response.status_code == 200
    # vérifie que tous les templates sont bien chargés dans la page résultante
    assertTemplateUsed(response, "base.html")
    assertTemplateUsed(response, "snippets/tableau_de_bord_menu.html")
    assertTemplateUsed(response, "vsme/exigence_de_publication.html")
    assert response.context["exigence_de_publication"] == exigence_de_publication
    assert f"Exercice {rapport_vsme.annee}" in response.content.decode()
    assert (
        f"Du 01/01/{rapport_vsme.annee} au 31/12/{rapport_vsme.annee}"
        in response.content.decode()
    )


@pytest.mark.parametrize("exigence_de_publication", EXIGENCES_DE_PUBLICATION.values())
def test_exigence_de_publication_vsme_avec_exercice_ne_correspondant_pas_a_une_annee_civile(
    exigence_de_publication, client, alice, rapport_vsme
):
    entreprise = rapport_vsme.entreprise
    entreprise.date_cloture_exercice = date(2020, 6, 30)
    entreprise.save()
    client.force_login(alice)

    url = EXIGENCE_DE_PUBLICATION_URL.format(
        vsme_id=rapport_vsme.id,
        exigence_de_publication_code=exigence_de_publication.code,
    )
    response = client.get(url)

    assert response.context["exigence_de_publication"] == exigence_de_publication
    assert (
        f"Exercice {rapport_vsme.annee - 1}-{rapport_vsme.annee}"
        in response.content.decode()
    )
    assert (
        f"Du 01/07/{rapport_vsme.annee - 1} au 30/06/{rapport_vsme.annee}"
        in response.content.decode()
    )


def test_exigence_de_publication_vsme_inexistante_retourne_une_404(
    client, alice, rapport_vsme
):
    client.force_login(alice)

    url = EXIGENCE_DE_PUBLICATION_URL.format(
        vsme_id=rapport_vsme.id, exigence_de_publication_code="Z1"
    )
    response = client.get(url)

    assert response.status_code == 404


def test_exigence_de_publication_vsme_d_un_rapport_inexistant_retourne_une_404(
    client, alice
):
    client.force_login(alice)

    url = EXIGENCE_DE_PUBLICATION_URL.format(
        vsme_id="yolo", exigence_de_publication_code="B1"
    )

    response = client.get(url)

    assert response.status_code == 404


def test_exigence_de_publication_vsme_est_privee(client, bob, rapport_vsme):
    url = EXIGENCE_DE_PUBLICATION_URL.format(
        vsme_id=rapport_vsme.id, exigence_de_publication_code="B1"
    )
    response = client.get(url)

    assert response.status_code == 302
    connexion_url = reverse("users:login")
    assert response.url == f"{connexion_url}?next={url}"

    # Bob n'est pas rattaché à l'entreprise du rapport VSME
    client.force_login(bob)
    response = client.get(url)

    assert response.status_code == 403


# Partage du rapport VSME par e-mail


def test_categories_vsme_formulaire_de_partage_non_valide_en_get(
    client, entreprise_factory, alice
):
    entreprise = entreprise_factory(utilisateur=alice)
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.get(url)

    form = response.context["form"]
    assert not form.is_bound


def test_categories_vsme_partage_du_rapport_par_email(
    client, entreprise_factory, alice, mailoutbox, mocker
):
    entreprise = entreprise_factory(
        utilisateur=alice, date_cloture_exercice=date(2024, 6, 30)
    )
    mocker.patch.object(
        RapportVSME, "est_complet", new_callable=mocker.PropertyMock, return_value=True
    )
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.post(
        url,
        {"email": "destinataire@example.com"},
        follow=True,
    )

    assert response.status_code == 200
    assert response.redirect_chain == [(url, 302)]
    assert (
        "Le rapport VSME a été envoyé à destinataire@example.com."
        in response.content.decode("utf-8")
    )
    assert len(mailoutbox) == 1
    mail = mailoutbox[0]
    assert mail.from_email == settings.DEFAULT_FROM_EMAIL
    assert list(mail.to) == ["destinataire@example.com"]
    assert mail.template_id == settings.BREVO_PARTAGE_VSME_TEMPLATE
    assert mail.merge_global_data == {
        "nom_entreprise": entreprise.denomination,
        "exercice": "2023-2024",
    }
    assert len(mail.attachments) == 2
    nom_fichier, contenu, type_mime = mail.attachments[0]
    assert nom_fichier == "vsme.xlsx"
    assert type_mime == XLSX_CONTENT_TYPE
    load_workbook(BytesIO(contenu))
    nom_fichier, contenu, type_mime = mail.attachments[1]
    assert nom_fichier == "vsme.pptx"
    assert type_mime == PPTX_CONTENT_TYPE
    Presentation(BytesIO(contenu))


def test_categories_vsme_partage_du_rapport_echappe_l_email_dans_le_message(
    client, entreprise_factory, alice, mailoutbox, mocker
):
    entreprise = entreprise_factory(utilisateur=alice)
    mocker.patch.object(
        RapportVSME, "est_complet", new_callable=mocker.PropertyMock, return_value=True
    )
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.post(
        url, {"email": '"<script>alert(1)</script>"@example.com'}, follow=True
    )

    content = response.content.decode("utf-8")
    assert "<script>alert(1)</script>" not in content
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in content


def test_categories_vsme_partage_refuse_si_rapport_incomplet(
    client, entreprise_factory, alice, mailoutbox
):
    entreprise = entreprise_factory(utilisateur=alice)
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.post(url, {"email": "destinataire@example.com"})

    assert response.status_code == 403
    assert len(mailoutbox) == 0


def test_categories_vsme_partage_du_rapport_avec_email_invalide(
    client, entreprise_factory, alice, mailoutbox, mocker
):
    entreprise = entreprise_factory(utilisateur=alice)
    mocker.patch.object(
        RapportVSME, "est_complet", new_callable=mocker.PropertyMock, return_value=True
    )
    client.force_login(alice)

    url = INDICATEURS_VSME_URL.format(id_public=entreprise.id_public, annee=2024)
    response = client.post(url, {"email": "invalide"})

    assert response.status_code == 200
    assert response.context["form"].errors["email"]
    assert (
        "Le rapport VSME n'a pas été envoyé car l'adresse e-mail n'est pas valide."
        in response.content.decode("utf-8")
    )
    assert len(mailoutbox) == 0
